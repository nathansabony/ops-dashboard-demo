"""
Pick the tab a slot needs out of a multi-tab deal workbook.

Brokers send one workbook with everything in it — a rent roll, the actual P&L,
a pro forma, physical-condition notes (example workbooks,
2026-09). Dropped into the Rent Roll slot it should read the rent roll; dropped
into the T12 slot it should read the actual P&L — never the pro forma.

    python sheet_pick.py <workbook> <rent_roll|t12> <out.xlsx>

Classifies every tab by its NAME first (that is what a person labels it by) and
by its CONTENT when the name says nothing ("Sheet1"). Writes the chosen tab(s)
— values only, so formulas that pointed at other tabs keep their numbers — to
<out.xlsx>, which the normal rent-roll / T12 pipeline then reads as if it had
been sent on its own.

Emits JSON:
  {"sheets":[{"name","kind","hidden","reason"}],   every tab, classified
   "picked":["name",…],                            tabs written to out (may be [])
   "want":"rent_roll"|"t12",
   "split":true|false,                              false = nothing to narrow
   "out": "<path>"|null}

kind is one of rent_roll, retail_rent_roll, t12, proforma, other.
"""
import json
import re
import sys
from datetime import date, datetime

_REAL_STDOUT = sys.stdout
sys.stdout = sys.stderr

SCAN_ROWS = 60
SCAN_COLS = 30

# ── names ────────────────────────────────────────────────────────────────────
PROFORMA_RE = re.compile(r"pro[\s\-_]?forma|budget|projection|projected|forecast|underwrit|stabili[sz]ed|\bpf\b", re.I)
RETAIL_RE = re.compile(r"retail|commercial|\bstore|office", re.I)
RR_NAME_RE = re.compile(r"rent[\s\-_]*roll|\brr\b|unit[\s\-_]*(mix|list|detail)|tenant[\s\-_]*(list|roster|schedule)|lease[\s\-_]*(schedule|abstract)|\brent\s*schedule", re.I)
T12_NAME_RE = re.compile(r"t[\s\-_]*12|\b(12|twelve)[\s\-_]*(month|mo)s?\b|trailing|\bttm\b|p\s*&\s*l|\bpnl\b|\bp\s*and\s*l\b|profit|income[\s\-_]*(statement|stmt)|operating[\s\-_]*(statement|stmt|history)|\bfinancials?\b|\bactuals?\b|\boperations?\b|\bopex\b|historical", re.I)
OTHER_NAME_RE = re.compile(r"physical|condition|notes?\b|cap[\s\-_]*ex|photos?|summary|cover|assumption|comps?\b|contact|map|debt|loan|sources", re.I)

# ── content ──────────────────────────────────────────────────────────────────
RR_WORDS = [r"\bunit\b|unit\s*#|apt", r"tenant|resident|lessee", r"\brent\b", r"lease|move[\s\-]?in|expir"]
T12_WORDS = [r"income", r"expense", r"net operating|\bnoi\b", r"tax", r"insurance|utilit|repair"]


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return " ".join(str(v).split())


MONTH_RE = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s\-/.,']*(\d{2,4})?$", re.I)
TRUE_PROFORMA_RE = re.compile(r"pro[\s\-_]?forma|projection|projected|forecast|underwrit", re.I)


def _month_cells(r):
    """How many cells of a row are a month (text like 'Oct 2025' or a real date)."""
    n = 0
    for v in r:
        if isinstance(v, (datetime, date)) or (isinstance(v, str) and MONTH_RE.match(v.strip())):
            n += 1
    return n


def _content_kind(rows):
    blob = " \n".join(" | ".join(_txt(v) for v in r) for r in rows).lower()
    rr = sum(1 for w in RR_WORDS if re.search(w, blob))
    t12 = sum(1 for w in T12_WORDS if re.search(w, blob))
    title = " ".join(_txt(v) for r in rows[:6] for v in r)
    # A statement runs twelve month columns across; a rent roll has one row per unit and at most a date or two.
    # Yardi/RealPage statements say "Book = Accrual,Budget" (a ledger name) and name the tab "Report1" or "12 Month Statement":
    # those were read as a pro forma or a rent roll (three other properties T12s, 2026-10-22).
    if t12 >= 2 and any(_month_cells(r) >= 6 for r in rows):
        if TRUE_PROFORMA_RE.search(title):
            return "proforma", "title says pro forma/projection"
        return "t12", "twelve month columns with income / expense lines"
    if PROFORMA_RE.search(title):
        return "proforma", "title says pro forma/budget"
    if rr >= 3 and rr >= t12:
        return "rent_roll", "unit / tenant / rent columns"
    if t12 >= 3:
        return "t12", "income / expense / NOI lines"
    return "other", "no rent-roll or P&L shape found"


def classify(name, rows):
    """-> (kind, reason)."""
    if PROFORMA_RE.search(name):
        return "proforma", "tab name"
    if RR_NAME_RE.search(name):
        return ("retail_rent_roll" if RETAIL_RE.search(name) else "rent_roll"), "tab name"
    if T12_NAME_RE.search(name):
        return "t12", "tab name"
    if OTHER_NAME_RE.search(name):
        return "other", "tab name"
    kind, why = _content_kind(rows)
    if kind == "rent_roll":
        title = " ".join(_txt(v) for r in rows[:6] for v in r)
        if RETAIL_RE.search(title) and not re.search(r"apartment|residential", title, re.I):
            kind = "retail_rent_roll"
    return kind, why


# ── reading ──────────────────────────────────────────────────────────────────
def read_book(path):
    """[(name, hidden, rows)] with rows = full list of value tuples."""
    lower = str(path).lower()
    if not lower.endswith(".xls"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, data_only=True)
            out = []
            for ws in wb.worksheets:
                rows = [tuple(c.value for c in r) for r in ws.iter_rows()]
                out.append((ws.title, ws.sheet_state != "visible", rows))
            return out
        except Exception:
            if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
                raise
    import xlrd
    wb = xlrd.open_workbook(path)
    out = []
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                    except Exception:
                        pass
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                vals.append(v)
            rows.append(tuple(vals))
        out.append((sh.name, getattr(sh, "visibility", 0) != 0, rows))
    return out


def pick(sheets, want):
    """Names of the tabs to keep for `want`, best first."""
    cands = [s for s in sheets if s["kind"] == want and s["has_data"]]
    if want == "rent_roll":
        visible = [s for s in cands if not s["hidden"]]
        if visible:
            return [s["name"] for s in visible]
        if cands:
            return [s["name"] for s in cands]
        # a retail-only workbook still has a rent roll to read
        retail = [s for s in sheets if s["kind"] == "retail_rent_roll" and s["has_data"]]
        return [s["name"] for s in retail[:1]]
    # T12: exactly one statement. Prefer visible, then a name that says
    # actual/T12/trailing over a generic "Financials".
    def score(s):
        n = s["name"]
        return (not s["hidden"],
                bool(re.search(r"actual|t[\s\-_]*12|trailing|ttm|historical", n, re.I)),
                s["reason"] == "tab name")
    cands.sort(key=score, reverse=True)
    return [s["name"] for s in cands[:1]]


def write_subset(book, names, out_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    by_name = {n: rows for n, _, rows in book}
    for n in names:
        ws = wb.create_sheet(title=n[:31])
        for r in by_name[n]:
            ws.append(list(r))
    wb.save(out_path)


def main(path, want, out_path):
    want = {"rr": "rent_roll", "rent_roll": "rent_roll", "t12": "t12"}.get(want, want)
    book = read_book(path)
    sheets = []
    for name, hidden, rows in book:
        scan = [r[:SCAN_COLS] for r in rows[:SCAN_ROWS]]
        has_data = any(any(v not in (None, "") for v in r) for r in scan)
        kind, why = classify(name, scan) if has_data else ("other", "empty")
        sheets.append({"name": name, "kind": kind, "hidden": hidden, "reason": why, "has_data": has_data})
    result = {"sheets": [{k: s[k] for k in ("name", "kind", "hidden", "reason")} for s in sheets],
              "want": want, "picked": [], "split": False, "out": None}
    if len(book) > 1:
        names = pick(sheets, want)
        result["picked"] = names
        if names:
            write_subset(book, names, out_path)
            result["split"] = True
            result["out"] = out_path
    _REAL_STDOUT.write(json.dumps(result, default=str) + "\n")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        _REAL_STDOUT.write(json.dumps({"error": "Usage: python sheet_pick.py <workbook> <rent_roll|t12> <out.xlsx>"}) + "\n")
        sys.exit(1)
    try:
        main(sys.argv[1], sys.argv[2], sys.argv[3])
    except Exception as e:
        _REAL_STDOUT.write(json.dumps({"error": f"Could not read workbook tabs: {e}"}) + "\n")
