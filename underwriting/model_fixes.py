"""
Template fixes that apply to a model workbook, shared by the generator (every new model) and
repair_model.py (models already on the disk). Each is idempotent and returns a plain-language note.

Found by the 2026-10-22 audit of the live board:

  widen_solver      The Max Price Solver sheet tried 11 candidate cap rates, 5.50% to 8.00%, and interpolated.
                    A deal that already clears the IRR target at 5.50% (Willow Bay, IRR 21.2%) read "Below 5.50% -
                    expand range", and one that needs a cap above 8% (Southern Pines, Club at Sandy Shores,
                    Gardens of Pine Ranch, Maple Court) read "Above 8.00% - not achievable": five of twelve board
                    deals showed "no solve" because the grid was too narrow, not because there was no price.
                    The grid is now 19 caps, 3.00% to 12.00%, and the summary formulas span all of them.
  dashboard_irr_na  Dashboard A10 showed 0.0% when the model could not solve an IRR (the IRR row is
                    IFERROR(IRR(...),0)); the board card and the PDF said n/a. It now says n/a too.
  reassessment_policy  Station Heights (Ohio): Year-1 property tax jumped from $1.84M to $4.53M because the July template
                    always reset the assessment to the purchase price and applied the millage to all of it, while the
                    model's own county note said "a sale itself does not automatically trigger reassessment" (Example
                    County reappraises on a cycle). That one line took $2.7M off Year-1 NOI and made the deal unsolvable.
                    The newer template has a "Reassessment on Sale" input (Tax!B19); the July one did not, and the Max
                    Price Solver ignored it in both. It now exists everywhere, the solver uses it, and a county rule that
                    says a sale does not trigger reassessment sets it to 0%.
"""
import re
from copy import copy

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

SOLVER_FIRST_CAP = 0.03
SOLVER_STEP = 0.005
SOLVER_COLS = 19            # B..T


def widen_solver(wb):
    if "Max Price Solver" not in wb.sheetnames:
        return None
    ws = wb["Max Price Solver"]
    last = get_column_letter(1 + SOLVER_COLS)          # T
    if ws[f"{last}7"].value is not None:
        return None                                    # already wide
    src_col = 12                                       # L: the last of the original eleven
    for c in range(src_col + 1, 1 + SOLVER_COLS + 1):
        col = get_column_letter(c)
        for r in range(6, 38):
            s = ws.cell(r, src_col)
            d = ws.cell(r, c)
            if isinstance(s.value, str) and s.value.startswith("="):
                d.value = Translator(s.value, origin=s.coordinate).translate_formula(d.coordinate)
            d._style = copy(s._style)
        ws.column_dimensions[col].width = ws.column_dimensions["L"].width or ws.column_dimensions["B"].width or 13
    for i in range(SOLVER_COLS):
        c = 2 + i
        cap = round(SOLVER_FIRST_CAP + SOLVER_STEP * i, 4)
        ws.cell(7, c).value = cap
        ws.cell(6, c).value = f"Cap {cap:.2%}"
    lo, hi = f"{SOLVER_FIRST_CAP:.2%}", f"{SOLVER_FIRST_CAP + SOLVER_STEP * (SOLVER_COLS - 1):.2%}"
    irr, coc, caps = f"B34:{last}34", f"B35:{last}35", f"B7:{last}7"
    interp = lambda rng: (f"INDEX({rng},B40)+(INDEX({rng},B40+1)-INDEX({rng},B40))*(Inputs!$B$47-INDEX({irr},B40))"
                          f"/(INDEX({irr},B40+1)-INDEX({irr},B40))")
    ws["B40"].value = f'=COUNTIF({irr},"<"&Inputs!$B$47)'
    ws["B41"].value = (f'=IF(B40=0,"Below {lo} - not achievable in range",IF(B40>={SOLVER_COLS},"Above {hi} - not achievable in range",'
                       f"{interp(caps)}))")
    ws["B45"].value = f'=IF(ISNUMBER(B41),IF(B40=0,INDEX({coc},1),{interp(coc)}),"-")'
    return f"Max Price Solver now tries {SOLVER_COLS} cap rates ({lo} to {hi}); it tried 11 (5.50% to 8.00%) and showed no price for deals outside that."


def dashboard_irr_na(wb):
    if "Dashboard" not in wb.sheetnames:
        return None
    ws = wb["Dashboard"]
    want = "=IF(AND('Debt & Exit'!$B$25=0,'Debt & Exit'!$B$26<>1),\"n/a\",'Debt & Exit'!$B$25)"
    cur = str(ws["A10"].value or "")
    if cur == "='Exit Scenarios'!$C$14":
        # the July-era template reads the Exit Scenarios tab and shows the text "NM" when no IRR solves; same word as everywhere else
        ws["A10"].value = "=IF(ISNUMBER('Exit Scenarios'!$C$14),'Exit Scenarios'!$C$14,\"n/a\")"
        return "Dashboard levered IRR shows n/a, not NM, when the model could not solve an IRR."
    if cur != "='Debt & Exit'!$B$25":
        return None
    ws["A10"].value = want
    return "Dashboard levered IRR shows n/a, not 0.0%, when the model could not solve an IRR."


POLICY_LABEL = "Reassessment on Sale (% of Price)"
NO_RESET_RE = re.compile(r"(sale|sold)[^.]{0,60}\b(does not|doesn't|do not|not)\b[^.]{0,30}(trigger|reassess)", re.I)


def reassessment_policy(wb):
    if "Tax" not in wb.sheetnames:
        return None
    t = wb["Tax"]
    notes = []
    if str(t["A19"].value or "") != POLICY_LABEL:
        if str(t["B14"].value or "") != "='Debt & Exit'!$B$6":
            return None                                   # not a layout this fix knows
        t["A19"] = POLICY_LABEL
        t["B19"] = 1.0
        t["B19"].number_format = "0%"
        t["C19"] = ("INPUT. 100% = assessment resets to the purchase price on sale. 0% = no reset; Year 1 tax holds at today's "
                    "taxable value and only grows by the tax-growth assumption. Drives this tab AND every column of the Max Price Solver.")
        t["B14"] = "=$B$8+('Debt & Exit'!$B$6-$B$8)*$B$19"
        notes.append("Tax tab gained the Reassessment on Sale input (the July template always reset the assessment to the purchase price).")
    # the county rule the tax lookup recorded on this very tab
    rule = " ".join(str(t.cell(r, 2).value or "") for r in range(20, 32) if isinstance(t.cell(r, 2).value, str) and len(t.cell(r, 2).value) > 60)
    cur = t["B19"].value
    taxable = t["B8"].value
    if isinstance(cur, (int, float)) and cur == 1 and isinstance(taxable, (int, float)) and taxable > 0 and NO_RESET_RE.search(rule):
        t["B19"] = 0.0
        notes.append("Reassessment on Sale set to 0%: the county rule on the Tax tab says a sale does not trigger reassessment, so Year-1 tax "
                     "holds at today's taxable value instead of resetting to the purchase price.")
    # the Max Price Solver applies the same policy at every candidate price
    if "Max Price Solver" in wb.sheetnames:
        s = wb["Max Price Solver"]
        fixed = 0
        for c in range(2, 2 + SOLVER_COLS):
            col = get_column_letter(c)
            v = s.cell(13, c).value
            want = (f"=IF(OR(Tax!$B$8<=0,Tax!$B$9<=0),Tax!$B$13,(Tax!$B$8+({col}8-Tax!$B$8)*Tax!$B$19)*Tax!$B$9/1000+Tax!$B$10)")
            if isinstance(v, str) and v.startswith("=") and "Tax!$B$19" not in v and "Tax!$B$9/1000" in v:
                s.cell(13, c).value = want
                fixed += 1
        if fixed:
            s["A13"].value = "Year 1 Property Tax (per reassessment policy)"
            notes.append("The Max Price Solver now applies the Reassessment on Sale policy at every candidate price.")
    return " ".join(notes) or None


def apply_all(wb):
    return [n for n in (widen_solver(wb), dashboard_irr_na(wb), reassessment_policy(wb)) if n]
