"""
Post-generation finish: compute the workbook's numbers and extract the
Dashboard fast-facts + Checks summary as JSON for the dashboard tab.

Two engines, tried in order:
1. Excel COM full-rebuild recalc (local Windows server) — also persists cached
   values into the saved file, so the downloaded workbook opens pre-computed.
2. The `formulas` library (pure Python, works on Render/Linux) — validated
   2026-07-13 against the an approved reference workbook: reproduces every
   fast-fact to the penny. Values are computed for display only; the file
   keeps live formulas and computes on open in Excel.

If both engines fail, returns computed=false — the tab shows a download-only
state. No silent wrong numbers, ever.

Usage: python finish_model.py <model.xlsx>   -> JSON to stdout
"""
import json
import subprocess
import sys
import time
import warnings

warnings.filterwarnings("ignore")

from openpyxl import load_workbook


# ── engine 1: Excel COM ───────────────────────────────────────────────────────

def try_recalc_com(path):
    """Full-rebuild recalc via a fresh Excel COM instance. Returns (ok, error)."""
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        return False, "Excel/COM not available on this server"
    import win32com.client
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(path)
        xl.CalculateFullRebuild()
        wb.Save()
        try:
            wb.Close(SaveChanges=False)
        except Exception:
            pass  # transient RPC rejection after Save is survivable
        return True, None
    except Exception as e:
        return False, str(e)
    finally:
        try:
            xl.Quit()
        except Exception:
            pass
        del xl
        # zombie-Excel discipline: DispatchEx can orphan EXCEL.EXE holding the
        # file; kill strays so later reads/saves aren't clobbered.
        time.sleep(1.5)
        try:
            out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq EXCEL.EXE"],
                                 capture_output=True, text=True, timeout=15).stdout
            for line in out.splitlines():
                if line.startswith("EXCEL.EXE"):
                    subprocess.run(["taskkill", "/F", "/PID", line.split()[1]],
                                   capture_output=True, timeout=15)
        except Exception:
            pass


# ── engine 2: `formulas` library ─────────────────────────────────────────────

def try_compute_formulas(path):
    """Compute all formulas in-memory. Returns (getter, lookup, error):
    getter(sheet, row, col) -> value or None; lookup is the raw
    {(SHEET, "A1"): value} dict, reused by bake_values_into_xlsx()."""
    try:
        import formulas
    except ImportError:
        return None, None, "formulas library not installed"
    try:
        import re
        xl = formulas.ExcelModel().loads(path).finish()
        sol = xl.calculate()
        # normalize: "'[book.xlsx]SHEET NAME'!A7" -> ("SHEET NAME", "A7"),
        # case-insensitive, so key-format quirks can't null out the facts
        lookup = {}
        key_re = re.compile(r"^'\[[^\]]+\](.+)'!([A-Z]+\d+)$")
        for k, cell in sol.items():
            m = key_re.match(k)
            if not m:
                continue
            v = getattr(cell, "value", cell)
            try:
                while hasattr(v, "__len__") and not isinstance(v, str):
                    v = v[0]
                if not isinstance(v, str):
                    f = float(v)
                    v = int(f) if f == int(f) and abs(f) < 1e15 else f
            except Exception:
                continue
            lookup[(m.group(1).upper(), m.group(2))] = v

        def getter(sheet, row, col):
            coord = f"{chr(64 + col)}{row}" if col <= 26 else None
            return lookup.get((sheet.upper(), coord))

        return getter, lookup, None
    except Exception as e:
        return None, None, str(e)[:300]


# ── bake computed values into the saved file (so any viewer, not just Excel,
# shows the right numbers) ───────────────────────────────────────────────────

_XML_ESCAPE = {"&": "&amp;", "<": "&lt;", ">": "&gt;"}


def _xml_escape(s):
    return "".join(_XML_ESCAPE.get(ch, ch) for ch in s)


def bake_values_into_xlsx(path, lookup):
    """Inject <v> cached-value tags next to each formula's <f> tag, directly
    in the sheet XML — openpyxl has no API for writing formula+cache pairs,
    but this is exactly what Excel itself writes on save. Only touches cells
    that have a formula and (as generator.py always produces) no existing
    cached value; skips anything it isn't confident is safe to rewrite."""
    import re
    import shutil
    import zipfile
    from xml.etree import ElementTree as ET

    tmp = path + ".baking"
    with zipfile.ZipFile(path, "r") as zin:
        names = zin.namelist()
        wb_xml = zin.read("xl/workbook.xml")
        rels_xml = zin.read("xl/_rels/workbook.xml.rels")

        ns_wb = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                 "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        rid_by_name = {}
        for sheet_el in ET.fromstring(wb_xml).findall(".//m:sheets/m:sheet", ns_wb):
            rid_by_name[sheet_el.get("name")] = sheet_el.get(f"{{{ns_wb['r']}}}id")

        ns_rel = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        target_by_rid = {
            rel.get("Id"): rel.get("Target")
            for rel in ET.fromstring(rels_xml).findall(".//r:Relationship", ns_rel)
        }

        sheet_path_by_name = {}
        for name, rid in rid_by_name.items():
            target = target_by_rid.get(rid)
            if not target:
                continue
            # OPC targets may be part-relative ("worksheets/sheet1.xml", how
            # Excel saves) or package-absolute ("/xl/worksheets/sheet1.xml",
            # how raw openpyxl saves — this broke Render downloads until
            # 2026-07-13: "xl/" + lstrip("/") built xl/xl/... paths and every
            # sheet silently missed)
            if target.startswith("/"):
                sheet_path_by_name[name.upper()] = target.lstrip("/")
            else:
                sheet_path_by_name[name.upper()] = "xl/" + target

        # The (?<!/) matters: an empty styled cell is written SELF-CLOSING
        # (<c r="B20" s="52"/>), and without the lookbehind the open-tag group
        # swallowed it, so the match spanned the empty cell PLUS the next one.
        # repl then looked the value up under the empty cell's reference, found
        # nothing, and left the real formula cell unbaked. Net effect: every
        # formula cell sitting immediately right of a blank cell downloaded
        # blank — Debt & Exit C20:C23, Pro Forma D9/D34:D36, and so on.
        cell_re = re.compile(
            r'(<c r="([A-Z]+\d+)"[^>]*?(?<!/)>)((?:(?!</c>).)*?)(</c>)', re.DOTALL
        )

        empty_v_re = re.compile(r"<v\s*/>|<v>\s*</v>")

        def make_repl(sheet_upper, counter):
            def repl(m):
                open_tag, ref, inner, close_tag = m.group(1), m.group(2), m.group(3), m.group(4)
                if "<f" not in inner:
                    return m.group(0)  # not a formula cell
                # raw openpyxl writes an empty <v /> after each formula —
                # drop it so we don't end up with two value elements
                inner = empty_v_re.sub("", inner)
                if "<v>" in inner:
                    return m.group(0)  # already has a real cached value
                val = lookup.get((sheet_upper, ref))
                if val is None:
                    return m.group(0)
                # The formula engine hands back the text "empty" for a formula that points at a blank cell (Excel
                # shows 0). Written under an existing t="n" it made `<v>empty</v>` in a numeric cell - invalid XML that
                # neither openpyxl nor Excel can open: Station Heights and Harbor City Apartments (2026-08-02)
                # could not be opened at all. A blank reference is 0; any other text is typed as text.
                if isinstance(val, str) and val.strip().lower() == "empty":
                    val = 0
                if isinstance(val, str):
                    open_tag = re.sub(r'\s+t="[^"]*"', "", open_tag)
                    open_tag = open_tag[:-1] + ' t="str">'
                    inner = inner + f"<v>{_xml_escape(val)}</v>"
                else:
                    inner = inner + f"<v>{val}</v>"
                counter[0] += 1
                return open_tag + inner + close_tag
            return repl

        rewritten = {}
        baked = 0
        for sheet_name, arc_path in sheet_path_by_name.items():
            if arc_path not in names:
                continue
            xml = zin.read(arc_path).decode("utf-8")
            counter = [0]
            new_xml = cell_re.sub(make_repl(sheet_name, counter), xml)
            if counter[0]:
                rewritten[arc_path] = new_xml
                baked += counter[0]

        if not rewritten:
            return 0

        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = rewritten.get(item.filename)
                zout.writestr(item, data.encode("utf-8") if data else zin.read(item.filename))

    shutil.move(tmp, path)
    return baked


# ── fact extraction ──────────────────────────────────────────────────────────

def grab_facts(get):
    """Build the facts JSON via a cell getter(sheet, row, col) -> value."""
    def strip_pairs(val_row, label_row, cols=6):
        out = []
        for c in range(1, cols + 1):
            label = get("Dashboard", label_row, c)
            value = get("Dashboard", val_row, c)
            if label is None and value is None:
                continue
            out.append({"label": str(label or "").strip(), "value": value})
        return out

    facts = {
        "property": get("Dashboard", 1, 1),
        "address": get("Dashboard", 2, 1),
        "strip1": strip_pairs(4, 5),
        "strip2": strip_pairs(7, 8),
    }
    # The scenario table is found by its header, not a fixed row range: the
    # 8-year rebuild added two scenarios and three detail rows above it, and a
    # hardcoded range would silently scrape the detail rows instead.
    start = None
    for r in range(13, 60):
        label = get("Dashboard", r, 1)
        if label and str(label).strip().upper().startswith("RETURNS vs TARGETS".upper()):
            start = r + 2  # skip the column-header row
            break
    scenarios = []
    for r in range(start or 40, (start or 40) + 12):
        label = get("Dashboard", r, 1)
        v = get("Dashboard", r, 2)
        if not label or not str(label).strip():
            break  # blank row ends the table
        if v is None or isinstance(v, str):
            continue  # unscored row (e.g. the Yr2 CoC line) or a header
        scenarios.append({
            "scenario": str(label).strip(),
            "irr": v,
            "equity_multiple": get("Dashboard", r, 3),
        })
    facts["scenarios"] = scenarios
    facts["checks"] = {
        "ok": get("Checks", 2, 1),
        "needs_input": get("Checks", 2, 2),
        "review": get("Checks", 2, 4),
    }
    facts["computed"] = get("Dashboard", 7, 1) is not None
    return facts


def openpyxl_getter(path):
    wb = load_workbook(path, data_only=True)
    def get(sheet, row, col):
        return wb[sheet].cell(row, col).value
    return get


def main(path):
    result = {}
    ok, err = try_recalc_com(path)
    result["recalc"] = ok
    result["recalc_error"] = err
    if ok:
        result["engine"] = "excel"
        getter = openpyxl_getter(path)
    else:
        getter, lookup, ferr = try_compute_formulas(path)
        if getter:
            result["engine"] = "formulas"
            try:
                result["baked_cells"] = bake_values_into_xlsx(path, lookup)
            except Exception as e:
                result["bake_error"] = str(e)[:300]  # display still works; download may open blank until recalced
        else:
            result["engine"] = None
            result["formulas_error"] = ferr
            getter = openpyxl_getter(path)  # will show computed=false

    try:
        result["facts"] = grab_facts(getter)
    except Exception as e:
        result["facts"] = None
        result["facts_error"] = str(e)[:300]
    print(json.dumps(result, default=str))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(json.dumps({"error": "Usage: python finish_model.py <model.xlsx>"}))
        sys.exit(1)
    main(sys.argv[1])
