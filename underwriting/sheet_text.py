"""
Excel workbook (.xlsx or legacy .xls) -> plain text for Claude extraction.

Emits JSON: {"format": "xlsx"|"xls", "sheets": [{"name": ..., "rows": N}],
"text": "...", "chars": N}. Text is CSV-ish, one line per row, sheets
separated by "=== SHEET <name> ===" markers, empty rows collapsed. Dates are
ISO. This is deliberately dumb — all understanding happens in the model.
"""
import json
import sys
from datetime import date, datetime

MAX_COLS = 30
MAX_ROWS_PER_SHEET = 4000


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return str(v).replace("\n", " ").replace("\r", " ")


def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, MAX_ROWS_PER_SHEET),
                                    max_col=min(ws.max_column or 1, MAX_COLS), values_only=True):
                rows.append([_cell(v) for v in row])
            yield ws.title, rows
    finally:
        wb.close()


def read_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    for sh in wb.sheets():
        rows = []
        for r in range(min(sh.nrows, MAX_ROWS_PER_SHEET)):
            out = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                    except Exception:
                        pass
                out.append(_cell(v))
            rows.append(out)
        yield sh.name, rows


def extract(path):
    lower = str(path).lower()
    readers = [read_xls] if lower.endswith(".xls") else [read_xlsx, read_xls]
    last_err = None
    for reader in readers:
        try:
            sheets = list(reader(path))
            fmt = "xls" if reader is read_xls else "xlsx"
            break
        except Exception as e:  # wrong format for this reader — try the other
            last_err = e
            sheets = None
    if sheets is None:
        raise last_err

    parts, meta = [], []
    for name, rows in sheets:
        lines, blank_run = [], 0
        for row in rows:
            while row and row[-1] == "":
                row = row[:-1]
            if not row:
                blank_run += 1
                continue
            if blank_run and lines:
                lines.append("")  # collapse any run of blanks to one separator
            blank_run = 0
            lines.append(",".join(row))
        if lines:
            parts.append(f"=== SHEET {name} ===\n" + "\n".join(lines))
            meta.append({"name": name, "rows": len(lines)})
    text = "\n\n".join(parts)
    return {"format": fmt, "sheets": meta, "text": text, "chars": len(text)}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(json.dumps({"error": "Usage: python sheet_text.py <file.xlsx|file.xls>"}))
        sys.exit(1)
    try:
        print(json.dumps(extract(sys.argv[1])))
    except Exception as e:
        print(json.dumps({"error": f"Could not read workbook: {e}"}))
        sys.exit(0)
