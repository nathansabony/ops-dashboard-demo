"""
Finished underwriting model(s) -> the numbers a past-deal card promotes.

The owner, 2026-10-21: the Underwriting tab shows "the past deals that were run and
the numbers that we found based on the property", with the Excel model attached.
Those numbers already sit in every model the generator wrote: finish_model.py
bakes the computed values into the saved workbook, so the Dashboard tab reads
back without Excel. This reads them back; nothing is recalculated here.

    python deal_card.py <model.xlsx> [<model2.xlsx> ...]

Emits JSON {"cards": {"<path>": {...}}} - one entry per file. A file that is not
a house underwriting model (no Dashboard/Inputs tab) gets {"error": "..."}
instead, so importing a random workbook is refused, not silently shown.
A model that was never computed (no baked values) comes back computed=false;
the card then says so rather than showing zeros.
"""
import json
import re
import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
warnings.filterwarnings("ignore")

_REAL_STDOUT = sys.stdout
sys.stdout = sys.stderr

from openpyxl import load_workbook

# finish_model.grab_facts is the one place that knows where the Dashboard's
# fast-fact strips and scenario table live (it finds the scenario table by its
# header, because the 8-year rebuild moved it). Reuse it, don't copy it.
from finish_model import grab_facts


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def named_cell(wb, name):
    """Value of a one-cell defined name, or None."""
    try:
        dn = wb.defined_names[name]
        dest = next(iter(dn.destinations))
        sheet, ref = dest
        return wb[sheet][ref.replace("$", "")].value
    except Exception:
        return None


# A unit whose "resident" is Down, Model, Office, Employee... is not a vacant unit that could lease: it earns nothing and
# is not in the market. Pinecrest Village (2026-10-22): 359 of 996 units, so its occupancy of 39.8% is 62.2% of the units that can
# actually rent. The card and the report show both rather than letting the smaller one stand alone.
NON_REVENUE_NAME = re.compile(r"^(down|model|office|admin|employee|non[\s-]?rev|storage|maint|guest|corporate)", re.I)


def non_revenue_units(wb):
    try:
        ws = wb["Rent Roll"]
        hdr = next(r for r in range(14, 40) if str(ws.cell(r, 1).value or "").strip().lower() == "unit")
        name_col = next(c for c in range(2, 12) if "name" in str(ws.cell(hdr, c).value or "").lower())
        return sum(1 for r in range(hdr + 1, ws.max_row + 1)
                   if ws.cell(r, 1).value and NON_REVENUE_NAME.match(str(ws.cell(r, name_col).value or "").strip()))
    except Exception:
        return None


def uncomputed_formulas(path, wb_values):
    """Formula cells whose value was never saved with the file (a model written by openpyxl and never opened in Excel
    shows nothing in them until someone recalculates). Returns how many, so the board can say a model is not finished."""
    try:
        wb_f = load_workbook(path, data_only=False)
        n = 0
        for ws in wb_f.worksheets:
            if ws.title not in wb_values.sheetnames:
                continue
            wv = wb_values[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and wv[c.coordinate].value is None:
                        n += 1
        return n
    except Exception:
        return None


def card_for(path):
    wb = load_workbook(path, data_only=True)
    if "Dashboard" not in wb.sheetnames or "Inputs" not in wb.sheetnames:
        return {"error": "Not a house underwriting model (no Dashboard / Inputs tab)."}
    def get(sheet, row, col):
        return wb[sheet].cell(row, col).value
    facts = grab_facts(get)
    by_label = {}
    for strip in (facts.get("strip1") or []) + (facts.get("strip2") or []):
        by_label[strip["label"].upper()] = strip["value"]
    # the third strip (returns) is not in grab_facts: labels in row 11, values row 10
    for c in range(1, 7):
        label, value = get("Dashboard", 11, c), get("Dashboard", 10, c)
        if label:
            by_label[str(label).upper()] = value

    def pick(prefix):
        for k, v in by_label.items():
            if k.startswith(prefix):
                return v
        return None

    irr_target = num(named_cell(wb, "IrrTarget"))
    year_built = pick("YEAR BUILT")
    price = num(pick("PURCHASE PRICE"))
    card = {
        "computed": bool(facts.get("computed")) and price is not None,
        "name": str(get("Inputs", 9, 2) or facts.get("property") or "").strip(),
        "address": str(get("Inputs", 10, 2) or "").strip(),
        "year_built": year_built if isinstance(year_built, (int, float)) else (str(year_built) if year_built else None),
        "units": num(pick("UNITS")),
        "occupancy": num(pick("OCCUPANCY")),
        "avg_rent": num(pick("AVG RENT")),
        "price": price,
        "price_per_door": num(pick("PRICE PER DOOR")),
        "t12_noi": num(pick("T12 NOI")),
        "year1_noi": num(pick("YEAR 1 NOI")),
        "capex": num(pick("TOTAL CAPEX")),
        "max_price": num(pick("MAX PRICE")),
        "irr": num(pick("LEVERED IRR")),
        # the IRR tile of the earlier template shows the text "NM" when the model could not solve one (Station Heights,
        # equity multiple -0.45x); the current template shows 0%. Either way the equity was not recovered.
        "irr_failed": False,
        "equity_multiple": num(pick("EQUITY MULTIPLE")),
        "dscr": num(pick("YEAR 1 DSCR")),
        "coc": num(pick("YR2 CASH-ON-CASH")),      # The owner 2026-10-22: rank on IRR, then cash-on-cash
        "irr_target": irr_target,
        "hold_years": num(named_cell(wb, "HoldPeriodYears")),
        "scenarios": facts.get("scenarios") or [],
        "checks": facts.get("checks") or {},
    }
    irr_raw = pick("LEVERED IRR")
    em = card["equity_multiple"]
    card["irr_failed"] = bool((isinstance(irr_raw, str) and irr_raw.strip().upper() in ("NM", "N/M", "N/A", "#NUM!", "#DIV/0!", "#VALUE!"))
                              or (card["irr"] == 0 and em is not None and abs(em - 1) > 1e-9))
    card["non_revenue_units"] = non_revenue_units(wb)
    card["uncomputed_formulas"] = uncomputed_formulas(path, wb)
    m = re.search(r"LEVERED IRR \(SALE YR(\d+)\)", next((k for k in by_label if k.startswith("LEVERED IRR")), ""))
    if m and card["hold_years"] is None:
        card["hold_years"] = int(m.group(1))
    return card


def main(paths):
    cards = {}
    for p in paths:
        try:
            cards[p] = card_for(p)
        except Exception as e:
            cards[p] = {"error": f"Could not read this workbook: {e}"}
    _REAL_STDOUT.write(json.dumps({"cards": cards}, default=str) + "\n")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        _REAL_STDOUT.write(json.dumps({"error": "Usage: python deal_card.py <model.xlsx> ..."}) + "\n")
        sys.exit(1)
    main(sys.argv[1:])
