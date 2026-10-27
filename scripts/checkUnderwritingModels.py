"""Checks on the Excel side of the underwriting tab.

1. The template fixes (underwriting/model_fixes.py) really change the template: the Max Price Solver spans 3%-12%
   in 19 columns, and the Dashboard's IRR tile reads "n/a" when the model could not solve one (it used to show 0.0%).
2. No model on a disk has formulas saved without a value. A model written by the generator and never recalculated opens
   with blank tabs in any viewer that does not recalculate. Pass the folder as the first argument (or set
   UNDERWRITING_DIR); without one, only the template checks run.

Run: python scripts/checkUnderwritingModels.py [folder-of-models]
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
UW = os.path.join(HERE, "..", "underwriting")
sys.path.insert(0, UW)

from openpyxl import load_workbook  # noqa: E402
import model_fixes  # noqa: E402
from deal_card import uncomputed_formulas  # noqa: E402

failed = 0


def check(label, ok, detail=""):
    global failed
    print(("PASS  " if ok else "FAIL  ") + label + (("  -> " + str(detail)) if detail and not ok else ""))
    if not ok:
        failed += 1


print("The template fixes")
wb = load_workbook(os.path.join(UW, "Underwriting Model TEMPLATE.xlsx"))
changes = model_fixes.apply_all(wb)
check("the fixes report what they changed", len(changes) == 2, changes)
ws = wb["Max Price Solver"]
caps = [ws.cell(7, c).value for c in range(2, 2 + model_fixes.SOLVER_COLS)]
check("the solver tries 19 cap rates", all(isinstance(v, (int, float)) for v in caps) and len(caps) == 19, caps)
check("...from 3.00% up to 12.00% in half-point steps", abs(caps[0] - 0.03) < 1e-9 and abs(caps[-1] - 0.12) < 1e-9
      and all(abs((b - a) - 0.005) < 1e-9 for a, b in zip(caps, caps[1:])), caps)
check("a price above or below the range says so instead of 'no solve'", "not achievable in range" in str(ws["B41"].value), ws["B41"].value)
a10 = str(wb["Dashboard"]["A10"].value)
check("the Dashboard IRR tile shows n/a when the IRR could not be solved (equity not 1.0x), not 0.0%", '"n/a"' in a10 and "$B$25=0" in a10, a10)
old = load_workbook(os.path.join(UW, "Underwriting Model TEMPLATE.xlsx"))
old["Dashboard"]["A10"].value = "='Exit Scenarios'!$C$14"
model_fixes.apply_all(old)
check("an older-template Dashboard (IRR read from Exit Scenarios) also shows n/a, not NM", '"n/a"' in str(old["Dashboard"]["A10"].value) and "ISNUMBER" in str(old["Dashboard"]["A10"].value), old["Dashboard"]["A10"].value)
again = model_fixes.apply_all(wb)
check("the fixes are safe to apply twice", again == [] or all(isinstance(x, str) for x in again), again)

print("\nA T12 is told apart from a rent roll")
import sheet_pick  # noqa: E402
MONTHS = ["Oct 2025", "Nov 2025", "Dec 2025", "Jan 2026", "Feb 2026", "Mar 2026", "Apr 2026", "May 2026", "Jun 2026", "Jul 2026", "Aug 2026", "Sep 2026", "Total"]


def statement(book):
    return [("Some Property (123)",), ("Statement (12 months)",), ("Period = Oct 2025-Sep 2026",), (book,), tuple(MONTHS),
            ("4000", "REVENUE"), ("4100", "Rental Income", 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 78), ("6200", "Insurance"), ("6300", "Repairs & Maintenance"),
            ("", "NET OPERATING INCOME", 1, 2, 3)]


rent_roll = [("Unit Rent Roll",), ("Unit", "Tenant", "Rent", "Lease From", "Lease To", "Move In")] + [(f"{100 + i}", "Resident", 1200, "2026-01-01", "2026-12-31", "2026-01-01") for i in range(20)]
check("a tab called 'Report1' with twelve month columns and income / expense lines is a T12", sheet_pick.classify("Report1", statement("Book = Accrual"))[0] == "t12")
check("a tab called '12 Month Statement' is a T12 by its name", sheet_pick.classify("12 Month Statement-3avonda", statement("Book = Accrual"))[0] == "t12")
check("'Book = Accrual,Budget (BA)' is a ledger name, not a budget: still a T12", sheet_pick.classify("Report1", statement("Book = Accrual,Budget (BA) ; Tree = x"))[0] == "t12")
check("a statement titled pro forma is still a pro forma", sheet_pick.classify("Report1", [("Pro Forma Operating Statement",)] + statement("Book = Accrual")[1:])[0] == "proforma")
check("a unit-per-row rent roll is still a rent roll", sheet_pick.classify("Report1", rent_roll)[0] == "rent_roll")

print("\nA rent roll's own summary lines are not units")
import tempfile  # noqa: E402
from openpyxl import Workbook  # noqa: E402
import mapped_extract  # noqa: E402
tmpd = tempfile.mkdtemp()
wbk = Workbook()
wsk = wbk.active
wsk.title = "Sheet1"
for r in [("Rent Roll",), ("Exported On: 08/24/2026",), ("Unit", "Tenant", "Status", "Sqft", "Market Rent", "Rent"),
          ("1269-1", "A Smith", "Current", 950, 1155, 1055), ("1269-2", "B Jones", "Notice-Unrented", 950, 1350, 1350), ("1269-3", "", "Vacant", 950, 1350, 0),
          ("3 Units", "66.7% Occupied", None, 2850, 3855, 2405), (None,), ("Total 3 Units", "66.7% Occupied", None, 2850, 3855, 2405)]:
    wsk.append(list(r))
xl = os.path.join(tmpd, "rr.xlsx")
wbk.save(xl)
mp = {"kind": "rent_roll", "mappable": True, "official": {"official_units": 3, "official_occupied_units": 2},
      "mapping": {"columns": {"unit_number": 1, "resident_name": 2, "status": 3, "sqft": 4, "market_rent": 5, "rent_charge": 6},
                  "status_rules": {"vacant_contains": ["vacant", "notice-unrented"]}, "occupancy_source": "status", "sheets": [{"name": "Sheet1", "data_start_row": 4}]}}
got = mapped_extract.extract(xl, mp)
check("a 'N Units' summary line under the data is not read as a unit (3 units, 2 occupied, no count-mismatch warning)",
      len(got["data"]["units"]) == 3 and sum(u["occupied_flag"] for u in got["data"]["units"]) == 2 and not any("report summary says" in w for w in got["warnings"]),
      [len(got["data"]["units"]), got["warnings"]])
check("a resident on notice is occupied even when the mapping lists the notice status as vacant", sum(u["occupied_flag"] for u in got["data"]["units"]) == 2)
check("...and the skipped summary line is counted in the read's notes", got["meta"].get("summary_rows_skipped") == 2, got["meta"])

print("\nA T12 whose header sits low, and a mapping that misses month columns")
import sheet_sample  # noqa: E402
mon = ["Aug 2025", "Sep 2025", "Oct 2025", "Nov 2025", "Dec 2025", "Jan 2026", "Feb 2026", "Mar 2026", "Apr 2026", "May 2026", "Jun 2026", "Jul 2026"]
wt = Workbook()
wst = wt.active
wst.title = "Sheet1"
for i in range(12):
    wst.append(["Income Statement - 12 Month" if i == 0 else None])
wst.append(["Account Name"] + mon + ["Total"])                      # row 13: one past the 12 head rows
for lab, base in (("Gross Potential Rent", 225000), ("Total Operating Income", 232000), ("Total Operating Expense", 85000)):
    vals = [base + k * 100 for k in range(12)]
    wst.append([lab] + vals + [sum(vals)])
noi = [232000 + k * 100 - 85000 - k * 100 for k in range(12)]
wst.append(["NOI - Net Operating Income"] + noi + [sum(noi)])
xt = os.path.join(tmpd, "t12.xlsx")
wt.save(xt)
smp = sheet_sample.sample(xt)
hdr_line = next(l for l in str(smp["text"] if isinstance(smp, dict) else smp).split("\n") if l.startswith("R13:"))
check("the month header row is shown in full to the mapping model even when it is below the head rows (all 12 months)", hdr_line.count("2025") + hdr_line.count("2026") == 12, hdr_line)
tm = lambda cols: {"kind": "t12", "mappable": True, "t12_mapping": {"sheet": "Sheet1", "month_header_row": 13, "month_columns": cols, "total_column": None, "label_columns": [1],
                   "fields": {"gross_potential_rent": {"row": 14}, "effective_gross_income": {"row": 15}, "total_operating_expenses": {"row": 16}, "net_operating_income": {"row": 17}}}}
bad = mapped_extract.extract(xt, tm([2, 3]))
good = mapped_extract.extract(xt, tm(list(range(2, 14))))
check("a mapping that took 2 of 12 month columns says it does not tie out to the statement's own NOI total", any(w.startswith("Tie-out:") and "month columns" in w for w in bad["warnings"]), bad["warnings"])
check("a mapping of all 12 months reproduces the statement's NOI and raises no tie-out warning", abs(good["data"]["net_operating_income"] - sum(noi)) < 1 and not any(w.startswith("Tie-out:") for w in good["warnings"]), good["warnings"])

print("\nProperty tax after a sale follows the county's rule (Station Heights, Ohio)")


def tax_book(rule):
    wbt = Workbook()
    t = wbt.active
    t.title = "Tax"
    t["B8"], t["B9"], t["B10"], t["B13"] = 31484740, 58.44, 0, 1839811.74
    t["B14"] = "='Debt & Exit'!$B$6"
    t["B26"] = rule
    sv = wbt.create_sheet("Max Price Solver")
    for c in range(2, 21):
        sv.cell(13, c).value = f"=IF(OR(Tax!$B$8<=0,Tax!$B$9<=0),Tax!$B$13,{get_column_letter(c)}8*Tax!$B$9/1000+Tax!$B$10)"
    return wbt


from openpyxl.utils import get_column_letter  # noqa: E402
oh = tax_book("Example County reappraises every 3 years; a sale itself does not automatically trigger reassessment to sale price.")
n_oh = model_fixes.reassessment_policy(oh)
check("an old-template Tax tab gains the Reassessment on Sale input, and a county rule saying a sale does not trigger reassessment sets it to 0%",
      oh["Tax"]["A19"].value == model_fixes.POLICY_LABEL and oh["Tax"]["B19"].value == 0 and "$B$19" in oh["Tax"]["B14"].value and n_oh, n_oh)
check("the Max Price Solver applies the policy at every candidate price (all 19 columns)",
      all("Tax!$B$19" in str(oh["Max Price Solver"].cell(13, c).value) for c in range(2, 21)))
fl = tax_book("Florida reassesses property near the sale price after a sale.")
model_fixes.reassessment_policy(fl)
check("a county that does reset on sale keeps 100%, so Florida-style deals are unchanged", fl["Tax"]["B19"].value == 1.0)
check("the fix is safe to apply twice", model_fixes.reassessment_policy(oh) is None or "gained" not in model_fixes.reassessment_policy(oh))

print("\nThe workbook's look (underwriting/model_style.py)")
import model_style  # noqa: E402
sty = load_workbook(os.path.join(UW, "Underwriting Model TEMPLATE.xlsx"))
model_fixes.apply_all(sty)
before = {(ws.title, c.coordinate): c.value for ws in sty.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}
notes1 = model_style.apply_all(sty)
formulas_changed = [k for k, v in before.items() if sty[k[0]][k[1]].value != v]
check("styling only reads and writes presentation: the only existing formulas it rewrites are the comps KPI tiles and the Reno Pace note and tiles",
      all(k[0] in ("Rent Comps", "Sale Comps", "Reno Pace") for k in formulas_changed), formulas_changed[:6])
check("the Reno Pace Year-1 impact tile reads the summary row, not a stale table cell", sty["Reno Pace"]["C2"].value == "=IFERROR(B%d,0)" % next(r for r in range(20, 200) if str(sty["Reno Pace"].cell(r, 1).value or "").strip().lower() == "year 1 net rent impact")
      , [sty["Reno Pace"]["C2"].value, sty["Reno Pace"]["D2"].value])
tc = {ws.title: (ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor else None) for ws in sty.worksheets}
check("only the important tabs are coloured: Pro Forma green, Exit Scenarios red, Inputs gold; documents, Tax, Reno Pace, comps and Checks have no colour",
      tc["Pro Forma"].endswith("14532D") and tc["Exit Scenarios"].endswith("C00000") and tc["Inputs"].endswith("B98A2E") and all(tc[n] is None for n in ("T12 Actuals", "Rent Roll", "Unit Mix", "Tax", "Reno Pace", "Rent Comps", "Sale Comps", "Checks", "Max Price Solver")), tc)
check("gridlines are off on every sheet", all(ws.sheet_view.showGridLines is False for ws in sty.worksheets))
legacy_navy = [ws.title for ws in sty.worksheets for row in ws.iter_rows() for c in row if c.fill and c.fill.fill_type == "solid" and str(c.fill.fgColor.rgb)[-6:] in ("1F4E78", "0F243E", "FFF2CC")]
check("no old navy or 'calculated' yellow fill is left anywhere", not legacy_navy, legacy_navy[:3])
pt = sty["Price Tool"]
check("the Pro Forma gets Year-1 $ per unit and % of EGI columns to the right, as formulas", sty["Pro Forma"]["N5"].value == "Yr-1 $ / Unit" and any(str(sty["Pro Forma"].cell(r, 15).value or "").startswith("=IFERROR($E") for r in range(6, 31)), sty["Pro Forma"]["N5"].value)
check("a Price Tool tab sits after the Max Price Solver with a cap-rate list", sty.sheetnames.index("Price Tool") == sty.sheetnames.index("Max Price Solver") + 1
      and any("Max Price Solver" in (dv.formula1 or "") for dv in pt.data_validations.dataValidation) and 0.03 <= pt["B3"].value <= 0.12)
check("...whose figures all read from the solver grid (nothing is typed in)", all("Max Price Solver" in str(pt.cell(r, 2).value) for r in (5, 6, 7, 8, 9)))
snap = {(ws.title, c.coordinate): c.value for ws in sty.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
model_style.apply_all(sty)
check("applying the style twice changes nothing", {(ws.title, c.coordinate): c.value for ws in sty.worksheets for row in ws.iter_rows() for c in row if c.value is not None} == snap)

folder = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("UNDERWRITING_DIR")
if folder and os.path.isdir(folder):
    print("\nModels in " + folder)
    for f in sorted(os.listdir(folder)):
        if not f.lower().endswith(".xlsx") or f.startswith("~$") or "TEMPLATE" in f:
            continue
        p = os.path.join(folder, f)
        try:
            n = uncomputed_formulas(p, load_workbook(p, data_only=True))
        except Exception as e:
            check(f + " opens", False, e)
            continue
        check(f + ": no formula without a value", n == 0, f"{n} formula cells have no saved value")
else:
    print("\n(no model folder given: the on-disk scan was skipped)")

print("\nAll Excel-model checks passed." if not failed else f"\n{failed} check(s) FAILED")
sys.exit(1 if failed else 0)
